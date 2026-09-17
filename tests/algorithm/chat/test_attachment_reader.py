from importlib import import_module
from types import SimpleNamespace

ar = import_module('lazymind.chat.engine.attachment_reader')


def test_filter_chat_document_files():
    files = [
        '/data/a.png',
        '/data/report.pdf',
        '/data/note.docx',
        '/data/slide.pptx',
    ]
    assert ar.filter_chat_document_files(files) == [
        '/data/report.pdf',
        '/data/note.docx',
        '/data/slide.pptx',
    ]


def test_text_files_are_supported_chat_attachments():
    files = ['/data/notes.txt', '/data/draft.lmd', '/data/config.YAML', '/data/main.py']

    assert all(ar.is_chat_text_file(path) for path in files)
    assert all(ar.is_chat_attachment_file(path) for path in files)


def test_filter_chat_image_files():
    files = ['/data/a.png', '/data/report.pdf', '/data/b.JPEG']
    assert ar.filter_chat_image_files(files) == ['/data/a.png', '/data/b.JPEG']


def test_parse_attachment_content_routes_by_suffix(monkeypatch, tmp_path):
    pdf_path = tmp_path / 'demo.pdf'
    image_path = tmp_path / 'photo.png'
    pdf_path.write_text('dummy', encoding='utf-8')
    image_path.write_bytes(b'png')

    monkeypatch.setattr(ar, 'select_vision_model_role', lambda: 'vlm')
    monkeypatch.setattr(ar, 'read_chat_document_text', lambda path: f'parsed:{path}')
    monkeypatch.setattr(ar, 'extract_image_description', lambda path, **kwargs: 'blue sky photo')

    assert ar.parse_attachment_content(str(pdf_path)) == f'parsed:{pdf_path.resolve()}'
    assert ar.parse_attachment_content(str(image_path)) == 'blue sky photo'


def test_parse_attachment_content_reads_text_without_ocr(monkeypatch, tmp_path):
    text_path = tmp_path / 'config.yaml'
    text_path.write_text('name: 测试\nenabled: true\n', encoding='utf-8')
    monkeypatch.setattr(
        ar,
        'read_chat_document_text',
        lambda path: (_ for _ in ()).throw(AssertionError('text file must not use OCR')),
    )

    assert ar.parse_attachment_content(str(text_path)) == 'name: 测试\nenabled: true\n'


def test_read_chat_text_file_truncates_and_rejects_nul(tmp_path):
    long_text = tmp_path / 'long.log'
    long_text.write_text('abcdefgh', encoding='utf-8')
    binary_text = tmp_path / 'fake.txt'
    binary_text.write_bytes(b'plain\x00binary')

    truncated = ar.read_chat_text_file(str(long_text), max_chars=4)
    assert truncated.startswith('abcd')
    assert 'truncated after 4 characters' in truncated

    try:
        ar.read_chat_text_file(str(binary_text))
        assert False, 'expected ValueError'
    except ValueError as exc:
        assert 'NUL bytes' in str(exc)


def test_build_attachment_reference_prompt(monkeypatch, tmp_path):
    pdf_path = tmp_path / 'demo.pdf'
    image_path = tmp_path / 'photo.png'
    pdf_path.write_text('dummy', encoding='utf-8')
    image_path.write_bytes(b'png')

    monkeypatch.setattr(ar, 'select_vision_model_role', lambda: 'vlm')
    monkeypatch.setattr(
        ar,
        'read_chat_document_text',
        lambda path: f'parsed:{path}',
    )
    monkeypatch.setattr(
        ar,
        'extract_image_description',
        lambda path, **kwargs: 'blue sky photo',
    )

    prompt = ar.build_attachment_reference_prompt([str(pdf_path), str(image_path)])

    assert 'Attached File References' in prompt
    assert 'demo.pdf' in prompt
    assert f'parsed:{pdf_path.resolve()}' in prompt
    assert 'photo.png' in prompt
    assert 'blue sky photo' in prompt


def test_sanitize_for_prompt_template_escapes_placeholders():
    raw = '<|im_start|>user\n{query} /think\n{response}<|im_end|>'
    sanitized = ar._sanitize_for_prompt_template(raw)
    assert sanitized == '<|im_start|>user\n{ query } /think\n{ response }<|im_end|>'
    assert '{query}' not in sanitized
    assert '{response}' not in sanitized


def test_build_attachment_reference_prompt_sanitizes_document_body(monkeypatch, tmp_path):
    pdf_path = tmp_path / 'demo.pdf'
    pdf_path.write_text('dummy', encoding='utf-8')
    monkeypatch.setattr(
        ar,
        'read_chat_document_text',
        lambda path: 'template {query} then {response}',
    )

    prompt = ar.build_attachment_reference_prompt([str(pdf_path)])

    assert 'template { query } then { response }' in prompt
    assert '{query}' not in prompt
    assert '{response}' not in prompt


def test_parse_attachment_content_rejects_unsupported_suffix(tmp_path):
    bad = tmp_path / 'archive.zip'
    bad.write_bytes(b'zip')
    try:
        ar.parse_attachment_content(str(bad))
        raise AssertionError('expected ValueError')
    except ValueError as exc:
        assert 'Unsupported attachment type' in str(exc)


def test_read_chat_document_text_joins_nodes(monkeypatch):
    def reader(_path):
        return [SimpleNamespace(text='line one'), SimpleNamespace(text='line two')]

    monkeypatch.setattr(ar, '_get_document_reader', lambda: reader)
    assert ar.read_chat_document_text('/tmp/demo.pdf') == 'line one\n\nline two'


def test_read_chat_document_text_reads_docx_locally_without_ocr(monkeypatch, tmp_path):
    from docx import Document

    path = tmp_path / 'requirements.docx'
    document = Document()
    document.add_paragraph('第一段要求')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '指标'
    table.cell(0, 1).text = '性能'
    document.add_paragraph('最后一段')
    document.save(path)
    monkeypatch.setattr(
        ar,
        '_get_document_reader',
        lambda: (_ for _ in ()).throw(AssertionError('DOCX must not use OCR when local parsing succeeds')),
    )

    body = ar.read_chat_document_text(str(path))

    assert body.splitlines() == ['第一段要求', '指标\t性能', '最后一段']


def test_read_chat_document_text_falls_back_to_ocr_for_invalid_docx(monkeypatch, tmp_path):
    path = tmp_path / 'broken.docx'
    path.write_bytes(b'not-a-docx')

    def reader(_path):
        return [SimpleNamespace(text='OCR fallback')]

    monkeypatch.setattr(ar, '_get_document_reader', lambda: reader)

    assert ar.read_chat_document_text(str(path)) == 'OCR fallback'


def test_read_chat_document_text_falls_back_to_ocr_for_image_only_docx(monkeypatch, tmp_path):
    from docx import Document

    path = tmp_path / 'image-only.docx'
    Document().save(path)
    monkeypatch.setattr(
        ar,
        '_get_document_reader',
        lambda: lambda _path: [SimpleNamespace(text='OCR image text')],
    )

    assert ar.read_chat_document_text(str(path)) == 'OCR image text'


def test_image_reader_reuses_llm_with_image_format(monkeypatch, tmp_path):
    from unittest.mock import Mock
    path = tmp_path / 'screen.png'
    path.write_bytes(b'png')
    monkeypatch.setattr(ar, 'select_vision_model_role', lambda: 'llm')
    model = Mock(return_value='visible screen text')
    factory = Mock(return_value=model)
    monkeypatch.setattr(ar, 'AutoModel', factory)
    assert ar.extract_image_description(str(path)) == 'visible screen text'
    factory.assert_called_once_with(model='llm', type='vlm')
    assert 'screen.png' in model.call_args.args[0]


def test_unavailable_image_is_explained_instead_of_silently_skipped(monkeypatch, tmp_path):
    from lazymind.vision_model import VisionModelUnavailable
    path = tmp_path / 'image.png'
    path.write_bytes(b'png')
    def unavailable():
        raise VisionModelUnavailable('请配置视觉模型后重试。')
    monkeypatch.setattr(ar, 'select_vision_model_role', unavailable)
    prompt = ar.build_attachment_reference_prompt([str(path)])
    assert 'Image was NOT read' in prompt
    assert '请配置视觉模型后重试' in prompt
    assert 'Do not infer its contents' in prompt
